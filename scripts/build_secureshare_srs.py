"""
Generate SecureShare IEEE 29148-Conformant SRS Word Document (.docx)
with cover page, professional tables, assets inventory, and complete functional/non-functional specifications.
"""

import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from docx_helpers import (
    setup_page_margins, add_header_footer, add_cover_page,
    add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_bullet_p, add_callout,
    format_table, set_cell_background, set_cell_margins
)

def build_secureshare_srs():
    doc = docx.Document()
    setup_page_margins(doc)
    add_header_footer(doc, "SecureShare")

    # Cover Page
    authors = [
        ("A. sushmitha", "CH.SC.U4CYS23057"),
        ("U. Tejaswi", "CH.SC.U4CYS23047")
    ]
    add_cover_page(
        doc,
        title="SecureShare",
        subtitle="Secure File Sharing Platform with Automated Security Validation",
        course="20CYS495 — Capstone Project",
        authors=authors
    )

    # Table of Contents
    add_heading_1(doc, "Table of Contents")
    toc_data = [
        ("1. Introduction", "1.1 Purpose, 1.2 Conventions, 1.3 Audience, 1.4 Scope, 1.5 Definitions"),
        ("2. Overall Description", "2.1 Perspective, 2.2 Functions, 2.3 User Classes, 2.4 Environment, 2.5 Constraints, 2.6 Dependencies"),
        ("3. Functional Requirements", "FR-01 to FR-12 Specification Matrix & 3.1 Traceability Matrix"),
        ("4. External Interface Requirements", "4.1 User Interfaces, 4.2 Software Interfaces, 4.3 Communication Interfaces"),
        ("5. Non-Functional Requirements", "Performance, Scalability, Availability, Usability, Maintainability, Portability, Reliability"),
        ("6. Security Requirements & Assets Inventory", "6.1 System Assets Inventory & Sensitivity Classification, 6.2 Application Security Requirements"),
        ("7. Data Requirements", "Core Managed Entities & Schema Descriptions"),
        ("8. Use Case Descriptions", "UC-1 to UC-7 Use Case Walkthroughs"),
        ("9. Acceptance Criteria", "Mandatory Automated & Security Verification Criteria"),
        ("10. Assumptions & Dependencies", "Operational Assumptions & Environmental Dependencies")
    ]
    
    t_toc = doc.add_table(rows=len(toc_data)+1, cols=2)
    t_toc.rows[0].cells[0].paragraphs[0].text = "Section Title"
    t_toc.rows[0].cells[1].paragraphs[0].text = "Scope & Covered Modules"
    for idx, (sec, desc) in enumerate(toc_data):
        t_toc.rows[idx+1].cells[0].paragraphs[0].text = sec
        t_toc.rows[idx+1].cells[1].paragraphs[0].text = desc
    format_table(t_toc, [2.5, 4.0])
    doc.add_page_break()

    # 1. Introduction
    add_heading_1(doc, "1. Introduction")
    
    add_heading_2(doc, "1.1 Purpose")
    add_body_p(doc, 
        "This document specifies the software requirements for SecureShare, a secure web application that lets "
        "users upload and share files while the platform automatically scans every uploaded file for security threats "
        "and continuously validates the application's own security through its CI/CD pipeline. It is prepared in "
        "conformance with the structure recommended by IEEE 29148."
    )
    
    add_heading_2(doc, "1.2 Document Conventions")
    add_body_p(doc,
        "Functional requirements are labeled FR-nn, and project objectives are labeled O1 through O5. "
        "Priority levels (High/Medium/Low) indicate relative implementation priority."
    )
    
    add_heading_2(doc, "1.3 Intended Audience")
    add_bullet_p(doc, "Project guide and review panel evaluating project deliverables.", "Project Guide & Review Panel: ")
    add_bullet_p(doc, "The developer(s) implementing SecureShare.", "Developers: ")
    add_bullet_p(doc, "Future contributors extending the scanning engines or CI/CD integrations.", "DevSecOps Contributors: ")

    add_heading_2(doc, "1.4 Project Scope")
    add_body_p(doc,
        "SecureShare allows registered users to upload files and share them securely with others. Every uploaded file "
        "is automatically validated for its declared type, scanned by a malware detection engine, and checked for "
        "dangerous scripts or executables before it is made available for download. Files can be shared through expiring, "
        "optionally password-protected links. In parallel, the application's own codebase is continuously validated: "
        "every commit and deployment triggers automated static analysis (SAST), dependency vulnerability scanning, "
        "and dynamic analysis (DAST) within the CI/CD pipeline, and administrators can review all findings — both "
        "file-level and pipeline-level — from a central security dashboard."
    )
    add_callout(doc, 
        "Out of scope for this phase: real-time collaborative file editing, mobile native applications, and automated "
        "remediation of source-code vulnerabilities (findings are surfaced for manual review rather than auto-patched).", 
        title="OUT OF SCOPE"
    )

    add_heading_2(doc, "1.5 Definitions, Acronyms, and Abbreviations")
    abbrev_data = [
        ("SAST", "Static Application Security Testing — analysis of source code for vulnerabilities without executing it."),
        ("DAST", "Dynamic Application Security Testing — analysis of a running application for vulnerabilities."),
        ("CI/CD", "Continuous Integration / Continuous Deployment — the automated build, test, and release pipeline."),
        ("MFA", "Multi-Factor Authentication."),
        ("RBAC", "Role-Based Access Control."),
        ("KYC", "Know Your Customer — identity-verification documents referenced as an example file type."),
        ("Magic Bytes", "The initial bytes of a file that identify its true format, used to detect disguised file types.")
    ]
    t_abb = doc.add_table(rows=len(abbrev_data)+1, cols=2)
    t_abb.rows[0].cells[0].paragraphs[0].text = "Term / Acronym"
    t_abb.rows[0].cells[1].paragraphs[0].text = "Definition & Technical Context"
    for idx, (a, d) in enumerate(abbrev_data):
        t_abb.rows[idx+1].cells[0].paragraphs[0].text = a
        t_abb.rows[idx+1].cells[1].paragraphs[0].text = d
    format_table(t_abb, [1.5, 5.0])

    # 2. Overall Description
    add_heading_1(doc, "2. Overall Description")
    
    add_heading_2(doc, "2.1 Product Perspective")
    add_body_p(doc,
        "SecureShare is a standalone web application intended for organizations that routinely handle sensitive files — "
        "such as HR departments, hospitals, banks, and universities — where an uploaded malicious file (malware, "
        "dangerous scripts, or executables) could compromise users or infrastructure. It is not a component of a larger product line."
    )

    add_heading_2(doc, "2.2 Product Functions")
    add_body_p(doc, "At a high level, SecureShare performs the following functions, derived from its core objectives (O1–O5):")
    add_bullet_p(doc, "Provide secure user management and role-based access control.", "O1: ")
    add_bullet_p(doc, "Detect and neutralize malicious or unsafe uploaded files automatically.", "O2: ")
    add_bullet_p(doc, "Enable secure, controlled file sharing with strong data protection.", "O3: ")
    add_bullet_p(doc, "Maintain a complete, tamper-evident audit trail of all activity.", "O4: ")
    add_bullet_p(doc, "Continuously validate the application's own security through the CI/CD pipeline.", "O5: ")

    add_heading_2(doc, "2.3 User Classes and Characteristics")
    user_classes = [
        ("Standard User", "Uploads, shares, and downloads files; has no visibility into other users' files unless explicitly shared with them."),
        ("Administrator", "Manages user accounts, reviews quarantined files and malware detections, and monitors CI/CD security findings via the dashboard."),
        ("CI/CD Pipeline (system actor)", "Automatically triggers SAST, dependency scanning, and DAST on every commit and deployment.")
    ]
    t_usr = doc.add_table(rows=len(user_classes)+1, cols=2)
    t_usr.rows[0].cells[0].paragraphs[0].text = "User Class / Actor"
    t_usr.rows[0].cells[1].paragraphs[0].text = "Characteristics & Access Privileges"
    for idx, (u, d) in enumerate(user_classes):
        t_usr.rows[idx+1].cells[0].paragraphs[0].text = u
        t_usr.rows[idx+1].cells[1].paragraphs[0].text = d
    format_table(t_usr, [2.2, 4.3])

    add_heading_2(doc, "2.4 Operating Environment")
    add_bullet_p(doc, "Web application accessible via any modern browser over HTTPS.")
    add_bullet_p(doc, "Backend server environment (e.g., Linux-based) with object/file storage.")
    add_bullet_p(doc, "Malware scanning engine (e.g., ClamAV or equivalent) integrated as a scanning service.")
    add_bullet_p(doc, "CI/CD platform (e.g., GitHub Actions, GitLab CI, or Jenkins) integrated with SAST, dependency-scanning, and DAST tools.")

    add_heading_2(doc, "2.5 Design and Implementation Constraints")
    add_bullet_p(doc, "Every uploaded file must pass malware and type-validation checks before it can be downloaded or shared.")
    add_bullet_p(doc, "All data in transit must use HTTPS/TLS; no unencrypted endpoints are permitted.")
    add_bullet_p(doc, "The CI/CD pipeline must block deployment when a critical-severity security finding is detected.")

    add_heading_2(doc, "2.6 Assumptions and Dependencies")
    add_bullet_p(doc, "A malware-scanning engine with regularly updated signatures is available to the system.")
    add_bullet_p(doc, "The CI/CD environment has access to configured SAST, dependency-scanning, and DAST tools.")
    add_bullet_p(doc, "Users access the platform using modern, HTTPS-capable browsers.")
    add_bullet_p(doc, "Underlying object storage guarantees durability of uploaded files.")

    doc.add_page_break()

    # 3. Functional Requirements
    add_heading_1(doc, "3. Functional Requirements")
    add_body_p(doc, "The following table specifies the functional requirements of SecureShare, grouped implicitly by the objective they satisfy.")

    fr_matrix = [
        ("FR-01", "User Registration & Authentication", "The system shall allow users to register and log in securely, with passwords stored using a strong hashing algorithm (bcrypt/Argon2) and support for multi-factor authentication.", "High"),
        ("FR-02", "Role-Based Access Control", "The system shall support at least two roles — Standard User and Administrator — with permissions enforced on every file and sharing operation.", "High"),
        ("FR-03", "File Upload", "The system shall allow authenticated users to upload files through the web interface, subject to a configurable maximum file size.", "High"),
        ("FR-04", "File Type & Signature Validation", "The system shall validate each uploaded file's declared MIME type against its actual file signature (magic bytes) and reject or quarantine files with mismatched or disallowed types.", "High"),
        ("FR-05", "Automated Malware Scanning", "The system shall automatically scan every uploaded file using a malware/antivirus scanning engine before the file is made available for sharing or download.", "High"),
        ("FR-06", "Dangerous Script & Executable Detection", "The system shall detect and flag uploaded executables, shell/batch scripts, and macro-enabled documents as high-risk content.", "High"),
        ("FR-07", "Quarantine & Notification", "The system shall automatically quarantine any file flagged as malicious or high-risk and notify the uploader and administrator of the detection.", "High"),
        ("FR-08", "Secure File Sharing", "The system shall allow users to generate shareable links with configurable expiry time, optional password protection, and download limits.", "High"),
        ("FR-09", "Encryption At Rest & In Transit", "The system shall encrypt stored files (AES-256 or equivalent) and enforce HTTPS/TLS for all client-server and file-transfer communication.", "High"),
        ("FR-10", "Audit Logging", "The system shall log all uploads, downloads, share actions, scan results, and administrative actions with user identity and timestamp.", "High"),
        ("FR-11", "CI/CD Security Pipeline Integration", "The system's build pipeline shall automatically run static application security testing (SAST), dependency vulnerability scanning, and dynamic application security testing (DAST) on every commit and deployment.", "High"),
        ("FR-12", "Security Dashboard & Reporting", "The system shall provide an administrator dashboard summarizing malware detections, quarantined files, and CI/CD pipeline security findings.", "Medium")
    ]

    t_fr = doc.add_table(rows=len(fr_matrix)+1, cols=4)
    t_fr.rows[0].cells[0].paragraphs[0].text = "ID"
    t_fr.rows[0].cells[1].paragraphs[0].text = "Title"
    t_fr.rows[0].cells[2].paragraphs[0].text = "Requirement Description"
    t_fr.rows[0].cells[3].paragraphs[0].text = "Priority"
    for idx, (f_id, title, desc, pri) in enumerate(fr_matrix):
        t_fr.rows[idx+1].cells[0].paragraphs[0].text = f_id
        t_fr.rows[idx+1].cells[1].paragraphs[0].text = title
        t_fr.rows[idx+1].cells[2].paragraphs[0].text = desc
        t_fr.rows[idx+1].cells[3].paragraphs[0].text = pri
    format_table(t_fr, [0.8, 1.8, 3.2, 0.7], col_alignments=["center", "left", "left", "center"])

    add_heading_2(doc, "3.1 Requirements Traceability")
    add_body_p(doc, "The table below traces each core objective to the functional requirements that realize it.")
    
    trace_data = [
        ("O1", "Provide secure user management and access control.", "FR-01, FR-02"),
        ("O2", "Detect and neutralize malicious or unsafe uploaded files automatically.", "FR-03, FR-04, FR-05, FR-06, FR-07"),
        ("O3", "Enable secure, controlled file sharing with strong data protection.", "FR-08, FR-09"),
        ("O4", "Maintain a complete, tamper-evident audit trail.", "FR-10"),
        ("O5", "Continuously validate the application's own security through the CI/CD pipeline.", "FR-11, FR-12")
    ]
    t_tr = doc.add_table(rows=len(trace_data)+1, cols=3)
    t_tr.rows[0].cells[0].paragraphs[0].text = "Objective"
    t_tr.rows[0].cells[1].paragraphs[0].text = "Description"
    t_tr.rows[0].cells[2].paragraphs[0].text = "Traced Requirements"
    for idx, (o, d, r) in enumerate(trace_data):
        t_tr.rows[idx+1].cells[0].paragraphs[0].text = o
        t_tr.rows[idx+1].cells[1].paragraphs[0].text = d
        t_tr.rows[idx+1].cells[2].paragraphs[0].text = r
    format_table(t_tr, [1.0, 4.0, 1.5], col_alignments=["center", "left", "left"])

    # 4. External Interface Requirements
    add_heading_1(doc, "4. External Interface Requirements")
    
    add_heading_2(doc, "4.1 User Interfaces")
    add_body_p(doc, 
        "SecureShare is operated through a responsive web interface providing screens for registration/login, "
        "file upload, a file/share management dashboard, and — for administrators — a security dashboard summarizing "
        "scan results and pipeline findings."
    )

    add_heading_2(doc, "4.2 Software Interfaces")
    add_bullet_p(doc, "Scans each uploaded file and returns a clean/malicious verdict (e.g., ClamAV or equivalent).", "Malware Scanning Engine: ")
    add_bullet_p(doc, "Stores uploaded files and their encrypted contents.", "Object/File Storage: ")
    add_bullet_p(doc, "Orchestrates automated SAST, dependency-vulnerability scanning, and DAST on each build/deploy.", "CI/CD Platform: ")
    add_bullet_p(doc, "Sends email/in-app alerts for quarantine events and critical pipeline findings.", "Notification Service: ")

    add_heading_2(doc, "4.3 Communication Interfaces")
    add_body_p(doc,
        "All client-server communication occurs over HTTPS. Internal service-to-service communication "
        "(e.g., application to scanning engine) uses REST APIs, secured within the deployment environment."
    )

    # 5. Non-Functional Requirements
    add_heading_1(doc, "5. Non-Functional Requirements")
    nfr_data = [
        ("Performance", "Malware scanning of a file up to 50 MB shall complete within 10 seconds under normal load."),
        ("Scalability", "The system shall support concurrent uploads and scans from at least 500 simultaneous users without service degradation."),
        ("Availability", "The platform shall maintain 99.5% uptime, excluding scheduled maintenance windows."),
        ("Usability", "Core actions — upload, share, download — shall each be completable within three user interactions from the dashboard."),
        ("Maintainability", "Malware signatures and CI/CD scanning tool rules shall be updatable without requiring application redeployment."),
        ("Portability", "The web application shall run on any modern browser and be deployable to any standard container-orchestration environment."),
        ("Reliability", "No file shall be made available for download before its malware scan has completed successfully.")
    ]
    t_nfr = doc.add_table(rows=len(nfr_data)+1, cols=2)
    t_nfr.rows[0].cells[0].paragraphs[0].text = "Attribute"
    t_nfr.rows[0].cells[1].paragraphs[0].text = "Requirement Statement"
    for idx, (a, r) in enumerate(nfr_data):
        t_nfr.rows[idx+1].cells[0].paragraphs[0].text = a
        t_nfr.rows[idx+1].cells[1].paragraphs[0].text = r
    format_table(t_nfr, [1.8, 4.7])

    doc.add_page_break()

    # 6. Security Requirements & Assets Inventory
    add_heading_1(doc, "6. Security Requirements & Assets Inventory")
    
    add_heading_2(doc, "6.1 System Assets Inventory & Sensitivity Classification")
    add_body_p(doc,
        "To satisfy comprehensive threat defense and IEEE 29148 security engineering rigor, the following primary "
        "assets are cataloged, classified by CIA impact (Confidentiality, Integrity, Availability), and mapped to security controls:"
    )

    asset_data = [
        ("AST-01", "User Accounts & Credentials", "Authentication", "C: High\nI: High\nA: Med", 
         "Brute-force attacks, credential stuffing, session hijacking.", 
         "bcrypt/Argon2 salted hashing, MFA enforcement, session timeout."),
        
        ("AST-02", "Uploaded File Payloads", "User Data", "C: High\nI: Critical\nA: High", 
         "Malware weaponization, data exfiltration, ransomware seeding.", 
         "AES-256 encryption at rest, pre-download ClamAV scanning, magic bytes validation."),
        
        ("AST-03", "Quarantine Store Vault", "Security Isolation", "C: High\nI: Critical\nA: High", 
         "Accidental execution of quarantined malicious binaries.", 
         "Strict isolation on separate unmapped storage path with zero execute permissions."),
        
        ("AST-04", "Share Links & Secrets", "Access Control", "C: High\nI: High\nA: Med", 
         "URL guessing, token enumeration, unauthorized link harvesting.", 
         "Cryptographically random high-entropy tokens, link expiration, password hashing."),
        
        ("AST-05", "Antivirus Signatures & Engine", "Security Service", "C: Med\nI: Critical\nA: Critical", 
         "Signature evasion, scanning service denial-of-service, engine tamper.", 
         "Automated hourly ClamAV signature updates, daemon isolation, fallback circuit breaker."),
        
        ("AST-06", "Tamper-Evident Audit Trail", "Compliance / Forensics", "C: Med\nI: Critical\nA: High", 
         "Log deletion by attacker to cover tracks, log injection.", 
         "Append-only storage, HMAC hashing, immutable actor/timestamp logging."),
        
        ("AST-07", "CI/CD Pipeline Secrets & Scanners", "DevSecOps Infrastructure", "C: Critical\nI: Critical\nA: Critical", 
         "Pipeline compromise, toxic dependency injection, secret leak in logs.", 
         "Encrypted secret manager, automated SAST, DAST, SCA dependency gating.")
    ]

    t_ast = doc.add_table(rows=len(asset_data)+1, cols=6)
    for i, h in enumerate(["Asset ID", "Asset Name", "Category", "CIA Rating", "Identified Threats", "Applied Security Controls"]):
        t_ast.rows[0].cells[i].paragraphs[0].text = h
    for idx, row in enumerate(asset_data):
        for c_idx, val in enumerate(row):
            t_ast.rows[idx+1].cells[c_idx].paragraphs[0].text = val
    format_table(t_ast, [0.7, 1.2, 1.0, 0.9, 1.3, 1.4], col_alignments=["center", "left", "left", "center", "left", "left"])

    add_heading_2(doc, "6.2 Application Security Requirements")
    add_bullet_p(doc, "Passwords shall be stored using a strong, salted hashing algorithm (bcrypt/Argon2); plaintext passwords shall never be logged or stored.")
    add_bullet_p(doc, "All communication shall be encrypted in transit using HTTPS/TLS; all stored files shall be encrypted at rest (AES-256).")
    add_bullet_p(doc, "Role-Based Access Control shall be enforced on every file, share, and administrative operation.")
    add_bullet_p(doc, "Every uploaded file shall be scanned for malware and validated by file signature before being made available for download or sharing.")
    add_bullet_p(doc, "Files flagged as malicious or high-risk shall be isolated in a quarantine store inaccessible to standard users.")
    add_bullet_p(doc, "User sessions shall expire after a period of inactivity, and all authentication events shall be logged.")
    add_bullet_p(doc, "All user input shall be validated and sanitized to prevent injection attacks (SQL injection, XSS, path traversal).")
    add_bullet_p(doc, "The CI/CD pipeline shall block deployment when a critical-severity vulnerability is detected by SAST, dependency scanning, or DAST.")
    add_bullet_p(doc, "Secrets and credentials used by the CI/CD pipeline shall be stored in a secure secrets manager, not in source code.")

    # 7. Data Requirements
    add_heading_1(doc, "7. Data Requirements")
    add_body_p(doc, "The core data entities managed by SecureShare are as follows:")
    
    data_entities = [
        ("User", "Registered account with credentials, role, and profile information."),
        ("File", "Metadata for an uploaded file, including name, size, MIME type, owner, and storage location."),
        ("FileVersion", "A specific version of a file, supporting version history where applicable."),
        ("ScanResult", "The outcome of a malware/security scan performed on a file, including verdict and engine details."),
        ("ShareLink", "A generated shareable link, including expiry time, password protection flag, and download limit."),
        ("AuditLog", "A record of a user or system action (upload, download, share, admin action) with timestamp and actor."),
        ("PipelineRun", "A record of a CI/CD build/deploy execution and its associated security scan results."),
        ("SecurityFinding", "A specific vulnerability or issue identified by SAST, DAST, or dependency scanning during a pipeline run.")
    ]
    t_de = doc.add_table(rows=len(data_entities)+1, cols=2)
    t_de.rows[0].cells[0].paragraphs[0].text = "Entity"
    t_de.rows[0].cells[1].paragraphs[0].text = "Description"
    for idx, (e, d) in enumerate(data_entities):
        t_de.rows[idx+1].cells[0].paragraphs[0].text = e
        t_de.rows[idx+1].cells[1].paragraphs[0].text = d
    format_table(t_de, [1.8, 4.7])

    doc.add_page_break()

    # 8. Use Case Descriptions
    add_heading_1(doc, "8. Use Case Descriptions")
    add_body_p(doc, "Primary actors: Standard User, Administrator, and the CI/CD Pipeline (system actor).")
    
    uc_list = [
        ("UC-1 Register / Login", "A user creates an account or authenticates securely, optionally using MFA."),
        ("UC-2 Upload File", "A user uploads a file, which is automatically validated and scanned before becoming available."),
        ("UC-3 Share File Securely", "A user generates a shareable link with an expiry time and optional password protection."),
        ("UC-4 Download Shared File", "A recipient accesses a shared file via a valid, unexpired link."),
        ("UC-5 Review Quarantined Files", "An administrator reviews files flagged as malicious or high-risk and confirms or releases them."),
        ("UC-6 Trigger CI/CD Security Scan", "On every commit or deployment, the pipeline automatically runs SAST, dependency scanning, and DAST."),
        ("UC-7 View Security Dashboard", "An administrator reviews a consolidated view of file-scan detections and CI/CD security findings.")
    ]
    t_uc = doc.add_table(rows=len(uc_list)+1, cols=2)
    t_uc.rows[0].cells[0].paragraphs[0].text = "Use Case ID & Title"
    t_uc.rows[0].cells[1].paragraphs[0].text = "Actor & Flow Summary"
    for idx, (u, f) in enumerate(uc_list):
        t_uc.rows[idx+1].cells[0].paragraphs[0].text = u
        t_uc.rows[idx+1].cells[1].paragraphs[0].text = f
    format_table(t_uc, [2.2, 4.3])

    # 9. Acceptance Criteria
    add_heading_1(doc, "9. Acceptance Criteria")
    add_bullet_p(doc, "Every uploaded file shall be scanned, and files containing known malware signatures shall be automatically quarantined.")
    add_bullet_p(doc, "Disguised executables and dangerous scripts shall be correctly identified via file-signature validation, regardless of file extension.")
    add_bullet_p(doc, "Shared files shall be accessible only through valid, unexpired share links, honoring any configured password protection and download limits.")
    add_bullet_p(doc, "A deployment containing a critical-severity SAST, dependency, or DAST finding shall be automatically blocked by the CI/CD pipeline.")
    add_bullet_p(doc, "All upload, download, share, and administrative actions shall appear in the audit log with correct timestamps and actor identity.")

    # 10. Assumptions
    add_heading_1(doc, "10. Assumptions & Dependencies")
    add_bullet_p(doc, "The malware-scanning engine's signature database is kept up to date.")
    add_bullet_p(doc, "The CI/CD environment has the necessary access and licenses for its configured SAST/DAST/dependency-scanning tools.")
    add_bullet_p(doc, "Users interact with the platform using modern, HTTPS-capable browsers.")
    add_bullet_p(doc, "Organizations deploying SecureShare (HR, healthcare, banking, education, etc.) provide their own compliance policies governing the specific document types they handle.")

    # Save
    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    out_rel = os.path.join(base_dir, "deliverables", "phase_02_requirements_engineering", "SecureShare_SRS_IEEE_29148.docx")
    doc.save(out_rel)
    print(f"[+] Saved SecureShare SRS to: {out_rel}")

    # Copy to Desktop
    user_home = os.environ.get("USERPROFILE") or os.environ.get("HOME")
    if user_home:
        desktop_dest = os.path.join(user_home, "Desktop", "SecureShare_SRS_IEEE_29148.docx")
        try:
            shutil.copyfile(out_rel, desktop_dest)
            print(f"[+] Copied SecureShare SRS to Desktop: {desktop_dest}")
        except Exception as e:
            print(f"[!] Desktop copy error: {e}")

if __name__ == "__main__":
    build_secureshare_srs()
