"""
TripMate & SecureShare IEEE 29148 SRS Document Generator
Produces publication-grade .docx files with cover page, professional tables,
assets inventory, and complete functional/non-functional specifications.
"""

import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    """Sets background shading of a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding (in twips) of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Sets subtle table borders."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_table(table, col_widths, col_alignments=None):
    """Formats a table with header styling, alternating rows, and widths."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, cell in enumerate(hdr_cells):
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
        cell.width = Inches(col_widths[i])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_alignments and col_alignments[i] == "center" else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                
    # Data rows
    for r_idx, row in enumerate(table.rows[1:]):
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for i, cell in enumerate(row.cells):
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            cell.width = Inches(col_widths[i])
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_alignments and col_alignments[i] == "center" else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(30, 41, 59)

def add_callout(doc, text, title="NOTE"):
    """Adds a stylish callout blockquote box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=200, right=160)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor(37, 99, 235)
    
    r_text = p.add_run(text)
    r_text.italic = True
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def setup_page_margins(doc):
    """Sets standard 1-inch margins."""
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

def add_header_footer(doc, title_text):
    """Sets up subtle headers and footers."""
    for section in doc.sections:
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"IEEE 29148 SRS — {title_text} | 20CYS495 Capstone Project")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun = fp.add_run("Amrita Vishwa Vidyapeetham — Dept. of Cybersecurity Systems & Networks")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

def add_cover_page(doc, title, subtitle, course, authors):
    """Creates a prestigious academic cover page."""
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(40)
    p_pre.paragraph_format.space_after = Pt(10)
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_pre = p_pre.add_run("SOFTWARE REQUIREMENTS SPECIFICATION")
    r_pre.font.name = "Calibri"
    r_pre.font.size = Pt(14)
    r_pre.font.bold = True
    r_pre.font.color.rgb = RGBColor(100, 116, 139)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(title)
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(30)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(4)
    p_sub.paragraph_format.space_after = Pt(24)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(51, 65, 85)
    
    # Divider line
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("—" * 38)
    r_div.font.color.rgb = RGBColor(203, 213, 225)
    
    # Standard badge
    p_std = doc.add_paragraph()
    p_std.paragraph_format.space_before = Pt(14)
    p_std.paragraph_format.space_after = Pt(18)
    p_std.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_std = p_std.add_run("IEEE 29148-Conformant SRS Specification")
    r_std.font.name = "Calibri"
    r_std.font.size = Pt(11)
    r_std.font.bold = True
    r_std.font.color.rgb = RGBColor(37, 99, 235)
    
    # Course
    p_course = doc.add_paragraph()
    p_course.paragraph_format.space_before = Pt(10)
    p_course.paragraph_format.space_after = Pt(40)
    p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_course = p_course.add_run(course)
    r_course.font.name = "Calibri"
    r_course.font.size = Pt(12)
    r_course.font.bold = True
    r_course.font.color.rgb = RGBColor(30, 41, 59)
    
    # Authors Box
    auth_table = doc.add_table(rows=len(authors)+1, cols=2)
    auth_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    auth_table.cell(0, 0).width = Inches(3.2)
    auth_table.cell(0, 1).width = Inches(2.8)
    
    # Header of authors table
    set_cell_background(auth_table.cell(0, 0), "1E3A8A")
    set_cell_background(auth_table.cell(0, 1), "1E3A8A")
    set_cell_margins(auth_table.cell(0, 0), top=100, bottom=100, left=140, right=140)
    set_cell_margins(auth_table.cell(0, 1), top=100, bottom=100, left=140, right=140)
    
    p1 = auth_table.cell(0, 0).paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = p1.add_run("Student Contributor")
    r1.bold = True
    r1.font.name = "Calibri"
    r1.font.size = Pt(10)
    r1.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = auth_table.cell(0, 1).paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run("Register Number")
    r2.bold = True
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(255, 255, 255)
    
    for idx, (name, reg) in enumerate(authors):
        c0 = auth_table.cell(idx+1, 0)
        c1 = auth_table.cell(idx+1, 1)
        bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_margins(c0, top=90, bottom=90, left=140, right=140)
        set_cell_margins(c1, top=90, bottom=90, left=140, right=140)
        
        cp0 = c0.paragraphs[0]
        cp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cr0 = cp0.add_run(name)
        cr0.font.name = "Calibri"
        cr0.font.bold = True
        cr0.font.size = Pt(10)
        cr0.font.color.rgb = RGBColor(30, 41, 59)
        
        cp1 = c1.paragraphs[0]
        cp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        cr1 = cp1.add_run(reg)
        cr1.font.name = "Calibri"
        cr1.font.size = Pt(10)
        cr1.font.color.rgb = RGBColor(37, 99, 235)
        
    set_table_borders(auth_table)
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(45)
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("Academic Year: 2025–2026 | Document Release: v1.0 Final")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(148, 163, 184)
    
    doc.add_page_break()

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138) # Navy
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(37, 99, 235) # Blue
    return h

def add_heading_3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(51, 65, 85) # Slate
    return h

def add_body_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Calibri"
        rb.font.size = Pt(10)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Calibri"
        rb.font.size = Pt(10)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

print("Helper definitions loaded successfully.")
